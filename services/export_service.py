from database.db import get_db_connection, get_dict_cursor
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from datetime import datetime


def export_students_excel():
    """Saare students ka Excel file banao"""
    conn = get_db_connection()
    cursor = get_dict_cursor(conn)
    cursor.execute("""
        SELECT 
            s.id, s.roll_number, u.full_name as name, u.email,
            c.name as class_name, sec.name as section_name
        FROM students s
        LEFT JOIN users u ON u.id = s.user_id
        LEFT JOIN classes c ON c.id = s.class_id
        LEFT JOIN sections sec ON sec.id = s.section_id
        ORDER BY s.roll_number
    """)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Students"

    # Header
    headers = ["ID", "Roll No", "Name", "Email", "Class", "Section"]
    ws.append(headers)

    # Header style
    header_fill = PatternFill(start_color="667EEA", end_color="667EEA", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data
    for r in rows:
        ws.append([
            r["id"], r["roll_number"], r["name"] or "",
            r["email"] or "", r["class_name"] or "", r["section_name"] or ""
        ])

    # Column widths
    widths = [8, 12, 25, 30, 15, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    # BytesIO mein save karo
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_attendance_excel(class_id: int = None):
    """Attendance ka Excel file banao (class-wise optional)"""
    conn = get_db_connection()
    cursor = get_dict_cursor(conn)

    if class_id:
        cursor.execute("""
            SELECT a.id, a.date, a.status, s.roll_number,
                   u.full_name as student_name, c.name as class_name
            FROM attendance a
            JOIN students s ON s.id = a.student_id
            LEFT JOIN users u ON u.id = s.user_id
            LEFT JOIN classes c ON c.id = s.class_id
            WHERE s.class_id = %s
            ORDER BY a.date DESC, s.roll_number
        """, (class_id,))
    else:
        cursor.execute("""
            SELECT a.id, a.date, a.status, s.roll_number,
                   u.full_name as student_name, c.name as class_name
            FROM attendance a
            JOIN students s ON s.id = a.student_id
            LEFT JOIN users u ON u.id = s.user_id
            LEFT JOIN classes c ON c.id = s.class_id
            ORDER BY a.date DESC, s.roll_number
        """)

    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    headers = ["ID", "Date", "Status", "Roll No", "Student Name", "Class"]
    ws.append(headers)

    header_fill = PatternFill(start_color="48BB78", end_color="48BB78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([
            r["id"], str(r["date"]), r["status"], r["roll_number"],
            r["student_name"] or "", r["class_name"] or ""
        ])

    widths = [8, 14, 12, 12, 25, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def export_fees_excel():
    """Fee payments ka Excel file banao"""
    conn = get_db_connection()
    cursor = get_dict_cursor(conn)
    cursor.execute("""
        SELECT fp.id, fp.amount, fp.monthly_fee, fp.yearly_fee,
               fp.payment_date, fp.payment_mod,
               s.roll_number, u.full_name as student_name
        FROM fee_payment fp
        JOIN students s ON s.id = fp.student_id
        LEFT JOIN users u ON u.id = s.user_id
        ORDER BY fp.payment_date DESC
    """)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Fees"

    headers = ["ID", "Roll No", "Student Name", "Amount", "Monthly", "Yearly", "Date", "Mode"]
    ws.append(headers)

    header_fill = PatternFill(start_color="F6AD55", end_color="F6AD55", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r in rows:
        ws.append([
            r["id"], r["roll_number"], r["student_name"] or "",
            r["amount"], r["monthly_fee"], r["yearly_fee"],
            str(r["payment_date"]), r["payment_mod"] or ""
        ])

    widths = [8, 12, 25, 12, 12, 12, 14, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output